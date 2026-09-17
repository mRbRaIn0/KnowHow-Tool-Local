"""Anhänge einer Chatnachricht: Ablage, Textextraktion, Übernahme in den Vault.

Angehängte Dateien landen zuerst in einem Zwischenbereich unter
data/uploads/<chat_id>/. Erst wenn das Modell sie ausdrücklich übernimmt,
werden sie in den Vault kopiert — das Original bleibt dabei immer erhalten.
"""
from __future__ import annotations

import base64
import io
import json
import logging
import re
import shutil
import time
import threading
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from .config import DATA_DIR
from .vault import IMAGE_EXT, create_unique_file, is_supported, kind_for

log = logging.getLogger(__name__)

UPLOAD_DIR = DATA_DIR / "uploads"

MAX_FILES = 50
MAX_FILE_BYTES = 40 * 1024 * 1024        # 40 MB je Datei
MAX_TOTAL_BYTES = 400 * 1024 * 1024      # 400 MB je Chat
MAX_EXTRACT_CHARS = 20000                # so viel Text bekommt das Modell je Datei
VISION_MAX_EDGE = 1400                   # Bilder werden dafür verkleinert
UPLOAD_TTL_SECONDS = 7 * 24 * 3600       # Zwischenablage nach einer Woche räumen


class AttachmentError(Exception):
    """Fehler beim Umgang mit Anhängen."""


class AnalysisIncomplete(AttachmentError):
    def __init__(self, message: str, partial: str = ""):
        super().__init__(message)
        self.partial = partial


def safe_filename(name: str) -> str:
    """Macht einen Dateinamen für Windows sicher, ohne ihn unkenntlich zu machen."""
    name = Path(name or "datei").name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", name).strip(". ")
    # Diese Zeichen sind unter Windows erlaubt, haben in Obsidian-WikiLinks
    # aber eine Steuerbedeutung (Überschrift, Blockanker oder Linksyntax).
    name = re.sub(r'[\[\]#^]', "-", name)
    if name.upper().split(".")[0] in {
        "CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)),
        *(f"LPT{i}" for i in range(1, 10)),
    }:
        name = f"_{name}"
    return name[:150] or "datei"


def chat_dir(chat_id: str) -> Path:
    """Zwischenbereich eines Chats. chat_id ist eine erzeugte Hex-ID."""
    if not re.fullmatch(r"[0-9a-f]{4,64}", chat_id or ""):
        raise AttachmentError("Ungültige Chat-Kennung.")
    target = UPLOAD_DIR / chat_id
    target.mkdir(parents=True, exist_ok=True)
    return target


def resolve(chat_id: str, name: str) -> Path:
    """Löst einen Anhangsnamen auf — ohne Ausbruch aus dem Zwischenbereich."""
    root = chat_dir(chat_id)
    target = (root / safe_filename(name)).resolve()
    if not str(target).startswith(str(root.resolve())):
        raise AttachmentError(f"Ungültiger Anhang: {name}")
    if not target.is_file():
        raise AttachmentError(f"Anhang '{name}' gibt es nicht.")
    return target


def store(chat_id: str, filename: str, data: bytes) -> Dict[str, Any]:
    """Legt eine hochgeladene Datei im Zwischenbereich ab."""
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(
            f"'{filename}' ist mit {len(data) // 1024 // 1024} MB zu groß "
            f"(erlaubt sind {MAX_FILE_BYTES // 1024 // 1024} MB)."
        )
    folder = chat_dir(chat_id)

    existing = [f for f in folder.iterdir() if f.is_file() and f.name not in INTERNE_DATEIEN]
    # Die Obergrenze gilt fuer die naechste Eingabe. Bereits verarbeitete
    # Anhaenge bleiben liegen, blockieren aber keinen neuen Upload.
    verwendet = load_used(chat_id)
    offen = [f for f in existing if f.name not in verwendet]
    if len(offen) >= MAX_FILES:
        raise AttachmentError(
            f"Es sind höchstens {MAX_FILES} Anhänge je Eingabe möglich."
        )
    if sum(f.stat().st_size for f in existing) + len(data) > MAX_TOTAL_BYTES:
        raise AttachmentError("Die Anhänge dieses Chats sind zusammen zu groß.")

    relative = create_unique_file(folder, safe_filename(filename), lambda stream: stream.write(data))
    return describe(folder / relative)


def describe(path: Path) -> Dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "kind": kind_for(path),
        "size": stat.st_size,
        "supported": is_supported(path),
        "modified": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
    }


def remove(chat_id: str, name: str) -> None:
    resolve(chat_id, name).unlink()


def clear(chat_id: str) -> None:
    try:
        shutil.rmtree(chat_dir(chat_id))
    except (OSError, AttachmentError):
        pass


def cleanup_old() -> int:
    """Räumt alte Zwischenablagen beim Start auf."""
    if not UPLOAD_DIR.is_dir():
        return 0
    entfernt = 0
    grenze = time.time() - UPLOAD_TTL_SECONDS
    for folder in UPLOAD_DIR.iterdir():
        try:
            if (folder.is_dir() and not (folder / CACHE_NAME).exists()
                    and folder.stat().st_mtime < grenze):
                shutil.rmtree(folder)
                entfernt += 1
        except OSError:
            continue
    return entfernt


# ------------------------------------------------------------ Textextraktion

def extract_text(path: Path, full: bool = False) -> Dict[str, Any]:
    """Holt Text aus einer Datei. Gibt immer ein Ergebnis zurück, nie None."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path, full)
        if suffix == ".docx":
            return _extract_docx(path, full)
        if suffix in IMAGE_EXT:
            return {"text": "", "hinweis": "Bilddatei — mit bild_ansehen betrachten."}
        return _extract_plain(path, full)
    except Exception as exc:  # defensiv: eine kaputte Datei darf nichts abbrechen
        log.warning("Extraktion von %s fehlgeschlagen: %s", path.name, exc)
        return {"text": "", "fehler": f"Datei konnte nicht gelesen werden: {exc}"}


def _cut(text: str, full: bool = False) -> Dict[str, Any]:
    gekuerzt = not full and len(text) > MAX_EXTRACT_CHARS
    return {"text": text[:MAX_EXTRACT_CHARS] if gekuerzt else text, "gekuerzt": gekuerzt}


def _extract_plain(path: Path, full: bool = False) -> Dict[str, Any]:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return _cut(data.decode(encoding), full)
        except UnicodeDecodeError:
            continue
    return {"text": "", "fehler": "Die Datei enthält keinen lesbaren Text."}


def _extract_pdf(path: Path, full: bool = False) -> Dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    seiten: List[str] = []
    for nummer, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            seiten.append(f"--- Seite {nummer} ---\n{text}")

    if not seiten:
        return {
            "text": "",
            "seiten": len(reader.pages),
            "hinweis": (
                "Dieses PDF enthält keinen auslesbaren Text — es ist ein Scan. "
                "Rufe bild_ansehen mit demselben Dateinamen auf; die Seiten werden "
                "dann als Bild ausgewertet."
            ),
        }
    result = _cut("\n\n".join(seiten), full)
    result["seiten"] = len(reader.pages)
    return result


def _extract_docx(path: Path, full: bool = False) -> Dict[str, Any]:
    import docx

    document = docx.Document(str(path))
    teile = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            zellen = [c.text.strip() for c in row.cells]
            if any(zellen):
                teile.append(" | ".join(zellen))
    return _cut("\n".join(teile), full)


# ------------------------------------------------------------------- Bilder

def pdf_page_images(path: Path, max_pages: int = 8) -> List[str]:
    """Rendert PDF-Seiten als Bilder — für Scans ohne Textebene.

    Genau der Fall 'Blatt mit dem iPhone abfotografiert und als PDF abgelegt'.
    """
    import pypdfium2 as pdfium

    seiten: List[str] = []
    document = pdfium.PdfDocument(str(path))
    try:
        for index in range(min(len(document), max_pages)):
            bild = document[index].render(scale=2).to_pil()
            if max(bild.size) > VISION_MAX_EDGE:
                from PIL import Image

                bild.thumbnail((VISION_MAX_EDGE, VISION_MAX_EDGE), Image.LANCZOS)
            buffer = io.BytesIO()
            bild.convert("RGB").save(buffer, format="JPEG", quality=85)
            seiten.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
    finally:
        document.close()
    return seiten


def page_count(path: Path) -> int:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(path)).pages)
    except Exception:
        return 0


def image_base64(path: Path, max_edge: int = VISION_MAX_EDGE) -> str:
    """Bild als base64 — verkleinert, damit die Bildanalyse zügig bleibt.

    Das Original auf der Platte bleibt unverändert.
    """
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.load()
            if max(image.size) > max_edge:
                image.thumbnail((max_edge, max_edge), Image.LANCZOS)
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=85)
            return base64.b64encode(buffer.getvalue()).decode("ascii")
    except Exception as exc:
        log.info("Bild %s wird unverkleinert übergeben (%s)", path.name, exc)
        return base64.b64encode(path.read_bytes()).decode("ascii")


CACHE_NAME = ".ausgewertet.json"
CACHE_VERSION = 3
_CACHE_LOCK = threading.RLock()


def _cache_path(chat_id: str) -> Path:
    return chat_dir(chat_id) / CACHE_NAME


def load_analysis(chat_id: str) -> dict:
    try:
        data = json.loads(_cache_path(chat_id).read_text(encoding="utf-8"))
        if isinstance(data, dict) and data.get("version") in (2, CACHE_VERSION):
            return {"items": data.get("items") or {}, "progress": data.get("progress") or {}}
    except (OSError, ValueError, AttachmentError):
        pass
    return {"items": {}, "progress": {}}


def load_cache(chat_id: str) -> Dict[str, str]:
    return load_analysis(chat_id)["items"]


def _save_analysis(chat_id: str, data: dict) -> None:
    target = _cache_path(chat_id)
    temporary = target.with_suffix(".tmp")
    data["version"] = CACHE_VERSION
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False)
        handle.flush()
        os.fsync(handle.fileno())
    temporary.replace(target)


def save_cache(chat_id: str, cache: Dict[str, str]) -> None:
    with _CACHE_LOCK:
        data = load_analysis(chat_id)
        data["items"].update(cache)
        _save_analysis(chat_id, data)


def checkpoint(chat_id: str, name: str, part: str = "", text: str = "",
               complete: bool = False, error: str = "") -> dict:
    """Commit every extracted unit atomically, including partial PDF progress."""
    with _CACHE_LOCK:
        data = load_analysis(chat_id)
        record = data["progress"].setdefault(name, {"parts": {}, "complete": False})
        if part:
            record["parts"].pop(part + ":partial", None)
            record["parts"][part] = text
        record.update(complete=complete, error=error)
        data["items"][name] = "\n\n".join(record["parts"].values())
        _save_analysis(chat_id, data)
        return record


def cached_excerpt(chat_id: str, name: str, offset: int = 0, limit: int = 7000) -> dict:
    data = load_analysis(chat_id)
    text = data["items"].get(name, "")
    offset = max(0, int(offset))
    limit = max(500, min(int(limit), 12000))
    end = min(len(text), offset + limit)
    status = data["progress"].get(name, {})
    return {"name": name, "text": text[offset:end], "gesamt_zeichen": len(text),
            "offset": offset, "naechster_offset": end if end < len(text) else None,
            "gekuerzt": end < len(text), "vollstaendig": bool(status.get("complete")),
            "hinweis": status.get("error", ""), "zwischengespeichert": True}


def pdf_text_pages(path: Path) -> List[str]:
    from pypdf import PdfReader
    return [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]


def pdf_page_image(path: Path, index: int) -> str:
    """Render just one page; never allocate all PDF bitmaps at once."""
    import pypdfium2 as pdfium
    from PIL import Image
    document = pdfium.PdfDocument(str(path))
    try:
        page = document[index]
        try:
            bitmap = page.render(scale=2)
            try:
                image = bitmap.to_pil()
                image.thumbnail((VISION_MAX_EDGE, VISION_MAX_EDGE), Image.Resampling.LANCZOS)
                buffer = io.BytesIO()
                image.convert("RGB").save(buffer, format="JPEG", quality=90)
                return base64.b64encode(buffer.getvalue()).decode("ascii")
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()


STATUS_NAME = ".verwendet.json"

# Verwaltungsdateien tauchen nie als Anhang auf.
INTERNE_DATEIEN = {CACHE_NAME, STATUS_NAME, ".ausgewertet.tmp"}


def _status_path(chat_id: str) -> Path:
    return chat_dir(chat_id) / STATUS_NAME


def load_used(chat_id: str) -> set:
    """Namen der Anhänge, die bereits an eine Nachricht übergeben wurden."""
    try:
        data = json.loads(_status_path(chat_id).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, AttachmentError):
        return set()
    namen = data.get("verwendet") if isinstance(data, dict) else None
    return set(namen) if isinstance(namen, list) else set()


def mark_used(chat_id: str, namen: Iterable[str]) -> None:
    """Merkt sich, dass diese Anhänge zu einer Eingabe gehört haben.

    Sie bleiben auf der Platte und über die Werkzeuge erreichbar, zählen aber
    nicht mehr zur nächsten Eingabe. So wird dieselbe Datei nicht mehrfach
    verarbeitet und der Fokus liegt auf dem, was neu dazukommt.
    """
    neu = {str(name) for name in namen if str(name).strip()}
    if not neu:
        return
    bestand = load_used(chat_id) | neu
    try:
        _status_path(chat_id).write_text(
            json.dumps({"verwendet": sorted(bestand)}, ensure_ascii=False),
            encoding="utf-8",
        )
    except (OSError, AttachmentError):
        log.warning("Anhang-Status konnte nicht geschrieben werden.")


def listing(chat_id: str, nur_offen: bool = False) -> List[Dict[str, Any]]:
    """Anhänge eines Chats.

    `nur_offen=True` liefert nur die noch nicht verwendeten — also genau das,
    was zur nächsten Eingabe gehört.
    """
    try:
        folder = chat_dir(chat_id)
    except AttachmentError:
        return []

    verwendet = load_used(chat_id) if nur_offen else set()
    eintraege = []
    for datei in folder.iterdir():
        if not datei.is_file() or datei.name in INTERNE_DATEIEN:
            continue
        if nur_offen and datei.name in verwendet:
            continue
        eintraege.append(describe(datei))
    return sorted(eintraege, key=lambda item: item["name"].lower())


def pending(chat_id: str) -> List[Dict[str, Any]]:
    """Kurzform für die Anhänge der nächsten Eingabe."""
    return listing(chat_id, nur_offen=True)


def summarise_for_prompt(items: List[Dict[str, Any]]) -> str:
    """Auflistung der Anhänge mit klarer Handlungsanweisung.

    Ohne die Anweisung antwortet das Modell gerne aus dem Gedächtnis, statt
    die Dateien tatsächlich anzusehen.
    """
    if not items:
        return ""

    zeilen = []
    for item in items:
        groesse = item["size"]
        lesbar = (f"{groesse / 1024 / 1024:.1f} MB" if groesse > 1024 * 1024
                  else f"{max(1, groesse // 1024)} KB")
        werkzeug = "bild_ansehen" if item["kind"] == "image" else "anhang_lesen"
        zeilen.append(
            f'- "{item["name"]}" ({KIND_LABELS.get(item["kind"], item["kind"])}, '
            f"{lesbar}) → {werkzeug}"
        )

    bilder = sum(1 for item in items if item["kind"] == "image")
    return (
        f"ANGEHÄNGTE DATEIEN ({len(items)}):\n" + "\n".join(zeilen) + "\n\n"
        "Diese Dateien liegen bereit. Du kennst ihren Inhalt NICHT — rate nicht "
        "und antworte nicht aus dem Gedächtnis. Rufe für jede Datei, die zur "
        "Aufgabe gehört, das angegebene Werkzeug auf, bevor du antwortest."
        + (f" {bilder} Bild(er) siehst du erst nach bild_ansehen." if bilder else "")
    )


KIND_LABELS = {
    "note": "Markdown", "text": "Textdatei", "code": "Quellcode",
    "doc": "Dokument", "image": "Bild", "other": "Datei",
}


def docx_images(path: Path) -> List[tuple[str, bytes]]:
    """Embedded Word images, including those without surrounding text."""
    import zipfile
    with zipfile.ZipFile(path) as archive:
        result = []
        for item in archive.infolist():
            if item.filename.startswith("word/media/") and not item.is_dir():
                if item.file_size > MAX_FILE_BYTES:
                    raise AttachmentError(f"Eingebettetes Bild zu groß: {item.filename}")
                result.append((item.filename, archive.read(item)))
        return result


def image_bytes_base64(data: bytes) -> str:
    from PIL import Image
    with Image.open(io.BytesIO(data)) as image:
        image.load()
        image.thumbnail((VISION_MAX_EDGE, VISION_MAX_EDGE), Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        image.convert("RGB").save(buffer, format="JPEG", quality=90)
        return base64.b64encode(buffer.getvalue()).decode("ascii")
