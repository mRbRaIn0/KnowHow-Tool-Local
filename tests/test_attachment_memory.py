import asyncio
import json
from unittest.mock import AsyncMock

import pytest

from backend import attachments
from backend.config import Profile
from backend.database import Database
from backend.routers import chat
from backend.tools import ToolRunner

CHAT = 'abcdef1234'

@pytest.fixture
def memory(tmp_path, monkeypatch):
    monkeypatch.setattr(attachments, 'UPLOAD_DIR', tmp_path / 'uploads')
    monkeypatch.setattr(attachments, 'image_base64', lambda path: path.name)
    return tmp_path


def test_ten_images_keep_full_notes_and_followup_uses_cache(memory):
    for index in range(10):
        attachments.store(CHAT, f'{index}.png', b'image')
    calls = []
    async def vision(prompt, image):
        calls.append(image)
        return image + ':' + 'Wort ' * 2000 + 'ENDMARKER'
    async def run():
        return [event async for event in chat._auswerten(CHAT, attachments.listing(CHAT), '', vision, None)]
    asyncio.run(run())
    assert len(calls) == 10
    assert all(text.endswith('ENDMARKER') for text in attachments.load_cache(CHAT).values())
    asyncio.run(run())
    assert len(calls) == 10
    runner = ToolRunner(memory, chat_id=CHAT)
    first = runner._anhang_lesen('0.png', limit=7000)
    second = runner._anhang_lesen('0.png', offset=first['naechster_offset'])
    assert second['text'].endswith('ENDMARKER')
    assert second['naechster_offset'] is None


def test_pdf_checkpoints_every_page_and_resumes_after_cancel(memory, monkeypatch):
    attachments.store(CHAT, 'scan.pdf', b'pdf')
    monkeypatch.setattr(attachments, 'pdf_text_pages', lambda path: ['TEXT'] * 12)
    monkeypatch.setattr(attachments, 'pdf_page_image', lambda path, index: index)
    calls = []
    async def vision(prompt, image):
        calls.append(image)
        if image == 3 and len(calls) == 4:
            raise asyncio.CancelledError()
        return f'VISUAL {image}'
    async def run():
        return [event async for event in chat._auswerten(CHAT, attachments.listing(CHAT), '', vision, None)]
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(run())
    saved = attachments.load_cache(CHAT)['scan.pdf']
    assert 'VISUAL 2' in saved
    assert not attachments.load_analysis(CHAT)['progress']['scan.pdf']['complete']
    asyncio.run(run())
    assert calls == [0, 1, 2, 3] + list(range(3, 12))
    assert 'VISUAL 11' in attachments.load_cache(CHAT)['scan.pdf']
    assert attachments.load_analysis(CHAT)['progress']['scan.pdf']['complete']


def test_plain_text_not_cut_at_20000_and_missing_vision_is_explicit(memory):
    attachments.store(CHAT, 'long.txt', ('A' * 25000 + 'TAIL').encode())
    attachments.store(CHAT, 'image.png', b'image')
    async def run():
        return [event async for event in chat._auswerten(CHAT, attachments.listing(CHAT), '', None, None)]
    events = asyncio.run(run())
    assert attachments.load_cache(CHAT)['long.txt'].endswith('TAIL')
    assert 'keine Bildauswertung' in events[-1]['progress']['image.png']['error']
    assert not events[-1]['progress']['image.png']['complete']


def test_thinking_only_has_visible_status_and_no_empty_fallback_file(memory, monkeypatch):
    profile = Profile(id='test', name='Test', vault={'path': str(memory)})
    db = Database(memory / 'chat.db')
    class Model:
        async def capabilities(self, model):
            return ['thinking', 'tools']
        async def chat_stream(self, *args, **kwargs):
            yield {'message': {'thinking': 'thinking'}, 'done': True, 'done_reason': 'length'}
    monkeypatch.setattr(chat, 'current_profile', lambda: profile)
    monkeypatch.setattr(chat, 'current_db', lambda profile=None: db)
    monkeypatch.setattr(chat, 'current_ollama', lambda profile=None: Model())
    monkeypatch.setattr(chat, 'hybrid_search', AsyncMock(return_value={}))
    async def run():
        work = db.create_chat('Work', 'local', 'vault')
        response = await chat.send_message(work['id'], chat.MessageRequest(content='Erstelle eine Notiz über Test.'))
        events = [json.loads(item[6:].strip()) async for item in response.body_iterator]
        assert events[-1]['type'] == 'done'
        assert 'keine abschließende Antwort' in events[-1]['content']
        assert events[-1]['changed_files'] == []
        assert db.list_messages(work['id'])[-1]['content']
    try:
        asyncio.run(run())
    finally:
        db.close()

def test_followup_receives_saved_source_notes_and_interruption_status(memory, monkeypatch):
    profile = Profile(id='test', name='Test', vault={'path': str(memory)})
    db = Database(memory / 'resume.db')
    captured = []
    class Model:
        async def capabilities(self, model):
            return ['tools']
        async def chat_stream(self, model, messages, **kwargs):
            captured.extend(messages)
            yield {'message': {'content': 'Die gespeicherten Arbeitsnotizen sind verfügbar.'}, 'done': True}
    monkeypatch.setattr(chat, 'current_profile', lambda: profile)
    monkeypatch.setattr(chat, 'current_db', lambda profile=None: db)
    monkeypatch.setattr(chat, 'current_ollama', lambda profile=None: Model())
    monkeypatch.setattr(chat, 'hybrid_search', AsyncMock(return_value={}))
    async def run():
        work = db.create_chat('Work', 'local', 'vault')
        attachments.store(work['id'], 'image.png', b'image')
        attachments.checkpoint(work['id'], 'image.png', 'image', 'Screenshot Kennung ZX-813', complete=True)
        attachments.mark_used(work['id'], ['image.png'])
        db.add_message(work['id'], 'user', 'Erfasse diesen Screenshot.')
        db.add_message(work['id'], 'assistant', 'Antwort unterbrochen. Arbeitsnotizen gespeichert.')
        response = await chat.send_message(work['id'], chat.MessageRequest(content='Was war die Kennung?'))
        events = [item async for item in response.body_iterator]
        assert events
        assert 'ZX-813' in json.dumps(captured)
        assert 'Antwort unterbrochen' in json.dumps(captured)
    try:
        asyncio.run(run())
    finally:
        db.close()


def test_cancel_after_tool_keeps_confirmed_action(memory, monkeypatch):
    profile = Profile(id='test', name='Test', vault={'path': str(memory)})
    db = Database(memory / 'cancel.db')
    calls = 0
    class Model:
        async def capabilities(self, model):
            return ['tools']
        async def chat_stream(self, *args, **kwargs):
            nonlocal calls
            calls += 1
            if calls > 1:
                raise asyncio.CancelledError()
            yield {'message': {'tool_calls': [{'function': {'name': 'notiz_erstellen',
                'arguments': {'pfad': 'Saved.md', 'inhalt': '# Saved\n\nSource content'}}}]}, 'done': True}
    monkeypatch.setattr(chat, 'current_profile', lambda: profile)
    monkeypatch.setattr(chat, 'current_db', lambda profile=None: db)
    monkeypatch.setattr(chat, 'current_ollama', lambda profile=None: Model())
    monkeypatch.setattr(chat, 'hybrid_search', AsyncMock(return_value={}))
    async def run():
        work = db.create_chat('Work', 'local', 'vault')
        response = await chat.send_message(work['id'], chat.MessageRequest(content='Erstelle eine Notiz.'))
        with pytest.raises(asyncio.CancelledError):
            async for _ in response.body_iterator:
                pass
        saved = db.list_messages(work['id'])[-1]
        assert 'Antwort unterbrochen' in saved['content']
        assert 'Saved.md' in saved['content']
        assert saved['sources'][0]['result']['erstellt'] == 'Saved.md'
    try:
        asyncio.run(run())
    finally:
        db.close()

def test_real_pdf_last_page_can_be_rendered_without_eight_page_limit(memory):
    from pypdf import PdfWriter
    import base64
    import io
    from PIL import Image
    path = memory / 'twelve.pdf'
    writer = PdfWriter()
    for _ in range(12):
        writer.add_blank_page(width=240, height=320)
    writer.write(path)
    assert len(attachments.pdf_text_pages(path)) == 12
    data = attachments.pdf_page_image(path, 11)
    with Image.open(io.BytesIO(base64.b64decode(data))) as image:
        assert image.size == (480, 640)

def test_docx_embedded_image_is_saved_with_text(memory):
    import docx
    from PIL import Image
    image = memory / 'embedded.png'
    Image.new('RGB', (32, 32), 'red').save(image)
    document = docx.Document()
    document.add_paragraph('Word source text')
    document.add_picture(str(image))
    path = memory / 'source.docx'
    document.save(path)
    attachments.store(CHAT, 'source.docx', path.read_bytes())
    async def vision(prompt, data):
        return 'Ein rotes Quadrat.'
    async def run():
        return [event async for event in chat._auswerten(CHAT, attachments.listing(CHAT), '', vision, None)]
    asyncio.run(run())
    text = attachments.load_cache(CHAT)['source.docx']
    assert 'Word source text' in text
    assert 'rotes Quadrat' in text
